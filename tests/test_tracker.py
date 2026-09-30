from pathlib import Path
from openpyxl import Workbook, load_workbook

from app.tracker import HEADERS, append_submission


def test_append_submission(tmp_path: Path):
    p = tmp_path / "tracker.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    for col, header in enumerate(HEADERS, 1):
        ws.cell(4, col).value = header
    wb.save(p)
    row = append_submission(p, {"platform":"Upwork","title":"Brand kit","url":"https://upwork.com/jobs/x","category":"Branding","budget_text":"$300","discovered_at":"2026-09-28T00:00:00Z"}, "manual test", 0)
    assert row == 5
    out = load_workbook(p, data_only=False)["Pipeline"]
    assert out.cell(5, 2).value == "Upwork"
    assert out.cell(5, 9).value == "Sent"
    assert p.with_suffix(".xlsx.bak").exists()
