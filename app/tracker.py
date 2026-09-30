from __future__ import annotations

import shutil
from copy import copy
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADERS = [
    "Date Found", "Platform", "Client / Project Title", "Category / Niche",
    "Budget (CAD/USD)", "Bids or Connects Used", "Date Applied",
    "Personalized Opening? (Y/N)", "Status", "Follow-Up Date",
    "Outcome / Notes", "Job Link",
]


def _first_empty_row(ws, start: int = 5, stop: int = 300) -> int:
    for row in range(start, stop + 1):
        if not ws.cell(row, 7).value and not ws.cell(row, 3).value:
            return row
    raise RuntimeError("Tracker Pipeline is full (rows 5-300). Extend the template before continuing.")


def append_submission(path: Path, job: dict[str, Any], notes: str = "", bid_units: int = 0) -> int:
    if not path.exists():
        raise FileNotFoundError(f"Tracker not found: {path}")
    backup = path.with_suffix(path.suffix + ".bak")
    shutil.copy2(path, backup)
    wb = load_workbook(path)
    if "Pipeline" not in wb.sheetnames:
        raise ValueError("Workbook must contain a 'Pipeline' sheet.")
    ws = wb["Pipeline"]
    actual = [ws.cell(4, c).value for c in range(1, 13)]
    if actual != HEADERS:
        raise ValueError("Pipeline columns do not match the expected Freelance_Application_Tracker schema.")
    row = _first_empty_row(ws)
    template_row = max(5, row - 1)
    for col in range(1, 13):
        src, dst = ws.cell(template_row, col), ws.cell(row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.font = copy(src.font)
        dst.protection = copy(src.protection)

    today = date.today()
    values = [
        job.get("discovered_at", today.isoformat())[:10], job["platform"], job["title"],
        job.get("category", "Other"), job.get("budget_text", ""), int(bid_units),
        today.isoformat(), "Y", "Sent", (today + timedelta(days=3)).isoformat(),
        notes, job["url"],
    ]
    for col, value in enumerate(values, 1):
        ws.cell(row, col).value = value
    wb.save(path)
    return row
