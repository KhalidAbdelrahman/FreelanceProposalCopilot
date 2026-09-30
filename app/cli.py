from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from .config import load_settings
from .connectors.browser_readonly import ReadOnlyBrowserDiscovery
from .connectors.m365_mail import M365MailReader
from .db import connect, session
from .services import _process_job, run_cycle


def import_samples(source: Path, target_dir: Path) -> int:
    target_dir.mkdir(parents=True, exist_ok=True)
    samples: list[str] = []
    if source.is_dir():
        for p in source.rglob("*"):
            if p.suffix.lower() in (".txt", ".md"):
                samples.append(p.read_text(encoding="utf-8"))
    elif source.suffix.lower() in (".txt", ".md"):
        samples.append(source.read_text(encoding="utf-8"))
    elif source.suffix.lower() == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                for key in ("proposal", "cover_letter", "application", "text"):
                    if row.get(key):
                        samples.append(row[key])
                        break
    elif source.suffix.lower() == ".xlsx":
        wb = load_workbook(source, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            headers = [str(v or "").strip().lower() for v in next(rows, [])]
            indexes = [i for i, h in enumerate(headers) if h in ("proposal", "cover letter", "application", "text")]
            for row in rows:
                for i in indexes:
                    if i < len(row) and row[i]:
                        samples.append(str(row[i]))
                        break
    else:
        raise ValueError("Supported inputs: directory of .txt/.md, .txt, .md, .csv, or .xlsx")
    clean = [s.strip() for s in samples if len(s.strip().split()) >= 20]
    for i, sample in enumerate(clean, 1):
        (target_dir / f"proposal_{i:04d}.txt").write_text(sample, encoding="utf-8")
    return len(clean)


def main() -> None:
    parser = argparse.ArgumentParser(description="Freelance Proposal Copilot")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("init-db")
    sub.add_parser("run-once")
    sub.add_parser("auth-m365")
    imp = sub.add_parser("import-proposals")
    imp.add_argument("source")
    browser = sub.add_parser("browser-scan")
    browser.add_argument("platform", choices=["upwork"])
    args = parser.parse_args()
    settings = load_settings()

    if args.command == "init-db":
        connect(settings.db_path).close()
        print(f"Initialized {settings.db_path}")
    elif args.command == "run-once":
        print(run_cycle(settings))
    elif args.command == "auth-m365":
        print("Authenticated" if M365MailReader().authenticate() else "Authentication failed")
    elif args.command == "import-proposals":
        count = import_samples(Path(args.source), settings.proposal_samples_dir)
        print(f"Imported {count} proposal samples into {settings.proposal_samples_dir}")
    elif args.command == "browser-scan":
        urls = settings.browser_saved_searches.get(args.platform, [])
        raw = ReadOnlyBrowserDiscovery(settings.enable_experimental_browser).discover_foreground(args.platform, urls)
        created = 0
        drafted = 0
        with session(settings.db_path) as conn:
            for job in raw:
                draft_id, is_new = _process_job(settings, conn, job)
                created += int(is_new)
                drafted += int(draft_id is not None)
        print(f"Recorded {created} new read-only discoveries and created {drafted} eligible drafts.")


if __name__ == "__main__":
    main()
